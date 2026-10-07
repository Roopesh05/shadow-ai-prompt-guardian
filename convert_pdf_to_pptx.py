"""
convert_pdf_to_pptx.py
======================
Converts 'Mini project.pdf' into a high-fidelity PowerPoint presentation (.pptx).
Preserves all 24 slides, logos, typography, titles, bullet structures,
architecture diagrams, and thank-you slide.
"""

import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_paths):
    prs = Presentation()
    
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_slide_layout = prs.slide_layouts[6] # completely blank layout
    
    logo_left_path = "extracted_pdf_assets/slide_1_img_1.jpeg"
    logo_right_path = "extracted_pdf_assets/slide_1_img_2.jpeg"
    diag_path = "extracted_pdf_assets/slide_12_img_1.jpeg"
    thank_you_path = "extracted_pdf_assets/slide_24_img_1.jpeg"
    
    def add_logos(slide):
        if os.path.exists(logo_left_path):
            slide.shapes.add_picture(logo_left_path, Inches(0.4), Inches(0.3), width=Inches(1.15), height=Inches(1.15))
        if os.path.exists(logo_right_path):
            slide.shapes.add_picture(logo_right_path, Inches(11.7), Inches(0.2), width=Inches(1.25), height=Inches(1.25))

    def add_title(slide, text, top=Inches(0.4), height=Inches(0.9), font_size=32):
        txBox = slide.shapes.add_textbox(Inches(1.8), top, Inches(9.733), height)
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)
        tf.margin_left = Inches(0)
        tf.margin_right = Inches(0)
        
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Times New Roman"
        p.font.size = Pt(font_size)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 0, 0)
        return txBox

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 1: TITLE SLIDE
    # ═════════════════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_slide_layout)
    add_logos(s1)
    
    # Title box (center)
    t_box = s1.shapes.add_textbox(Inches(1.8), Inches(1.8), Inches(9.733), Inches(1.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Mini Project Review"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = "Times New Roman"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(0, 0, 0)
    
    p2 = tf1.add_paragraph()
    p2.text = "Shadow AI Prompt Guardian"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Times New Roman"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(0, 0, 0)
    p2.space_before = Pt(14)
    
    # Bottom Left: Guide
    g_box = s1.shapes.add_textbox(Inches(0.6), Inches(5.4), Inches(4.5), Inches(1.5))
    gtf = g_box.text_frame
    gtf.word_wrap = True
    gp1 = gtf.paragraphs[0]
    gp1.text = "Under the Guidance of:"
    gp1.font.name = "Times New Roman"
    gp1.font.size = Pt(18)
    gp1.font.color.rgb = RGBColor(0, 0, 0)
    
    gp2 = gtf.add_paragraph()
    gp2.text = "Mr. NaveenRaj AP/CSE(IOT)"
    gp2.font.name = "Times New Roman"
    gp2.font.size = Pt(18)
    gp2.font.color.rgb = RGBColor(0, 0, 0)
    gp2.space_before = Pt(4)
    
    # Bottom Right: Team Members
    m_box = s1.shapes.add_textbox(Inches(8.2), Inches(4.5), Inches(4.8), Inches(2.5))
    mtf = m_box.text_frame
    mtf.word_wrap = True
    mp1 = mtf.paragraphs[0]
    mp1.text = "Team Members:"
    mp1.font.name = "Times New Roman"
    mp1.font.size = Pt(18)
    mp1.font.bold = True
    mp1.font.color.rgb = RGBColor(0, 0, 0)
    
    members = [
        "Ram Sakthi Sharumathi M (713324SB041)",
        "Roopesh S N (713324SB044)",
        "Sanjeevkumar P (713324SB046)",
        "Saravana Kumar M (713324SB048)"
    ]
    for mem in members:
        mp = mtf.add_paragraph()
        mp.text = mem
        mp.font.name = "Times New Roman"
        mp.font.size = Pt(17)
        mp.font.color.rgb = RGBColor(0, 0, 0)
        mp.space_before = Pt(4)

    # ═════════════════════════════════════════════════════════════════════════
    # HELPER: BULLET SLIDE GENERATOR
    # ═════════════════════════════════════════════════════════════════════════
    def create_bullet_slide(title, bullets, subheading=None, font_size=18, space_after=8):
        slide = prs.slides.add_slide(blank_slide_layout)
        add_logos(slide)
        add_title(slide, title)
        
        top_offset = Inches(1.5)
        content_box = slide.shapes.add_textbox(Inches(1.0), top_offset, Inches(11.333), Inches(5.5))
        tf = content_box.text_frame
        tf.word_wrap = True
        
        first = True
        if subheading:
            p_sub = tf.paragraphs[0]
            p_sub.text = subheading
            p_sub.font.name = "Times New Roman"
            p_sub.font.size = Pt(font_size + 2)
            p_sub.font.bold = True
            p_sub.font.color.rgb = RGBColor(0, 0, 0)
            p_sub.space_after = Pt(10)
            first = False
            
        for b_text in bullets:
            if first:
                p = tf.paragraphs[0]
                first = False
            else:
                p = tf.add_paragraph()
                
            p.font.name = "Times New Roman"
            p.font.size = Pt(font_size)
            p.space_after = Pt(space_after)
            p.line_spacing = 1.2
            
            # Formatting with bold prefix if it has "→" or ":"
            if "→" in b_text:
                parts = b_text.split("→", 1)
                r_bullet = p.add_run()
                r_bullet.text = "• "
                r_prefix = p.add_run()
                r_prefix.text = parts[0].strip().replace("•", "").strip() + " → "
                r_prefix.font.bold = True
                r_prefix.font.color.rgb = RGBColor(0, 0, 0)
                
                r_desc = p.add_run()
                r_desc.text = parts[1].strip()
                r_desc.font.color.rgb = RGBColor(30, 41, 59)
            elif ":" in b_text and not b_text.startswith("http"):
                parts = b_text.split(":", 1)
                r_bullet = p.add_run()
                r_bullet.text = "• "
                r_prefix = p.add_run()
                r_prefix.text = parts[0].strip().replace("•", "").strip() + ": "
                r_prefix.font.bold = True
                r_prefix.font.color.rgb = RGBColor(0, 0, 0)
                
                r_desc = p.add_run()
                r_desc.text = parts[1].strip()
                r_desc.font.color.rgb = RGBColor(30, 41, 59)
            else:
                clean_text = b_text.strip()
                if not clean_text.startswith("•"):
                    clean_text = "• " + clean_text
                r = p.add_run()
                r.text = clean_text
                r.font.color.rgb = RGBColor(30, 41, 59)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 2: PROBLEM STATEMENT (PARAGRAPHS)
    # ═════════════════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_logos(s2)
    add_title(s2, "Problem Statement")
    
    s2_box = s2.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.933), Inches(4.8))
    s2_tf = s2_box.text_frame
    s2_tf.word_wrap = True
    
    p2_1 = s2_tf.paragraphs[0]
    p2_1.text = "The increasing use of public AI chatbots creates a risk of accidentally sharing sensitive information. Users may unknowingly paste API keys, passwords, personal data, source code, or confidential business information into AI prompts. Once the prompt is submitted, this information may be transmitted to an external AI service."
    p2_1.font.name = "Times New Roman"
    p2_1.font.size = Pt(21)
    p2_1.line_spacing = 1.3
    p2_1.space_after = Pt(24)
    p2_1.font.color.rgb = RGBColor(0, 0, 0)
    
    p2_2 = s2_tf.add_paragraph()
    p2_2.text = "Traditional DLP solutions may require additional network infrastructure or centralized monitoring. Therefore, a lightweight client-side solution is needed to detect and prevent sensitive data leakage before the prompt is transmitted."
    p2_2.font.name = "Times New Roman"
    p2_2.font.size = Pt(21)
    p2_2.line_spacing = 1.3
    p2_2.font.color.rgb = RGBColor(0, 0, 0)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 3: ABSTRACT (PART 1)
    # ═════════════════════════════════════════════════════════════════════════
    s3_bullets = [
        "SENTINEL is a client-side Data Loss Prevention (DLP) browser extension.",
        "It protects sensitive information entered into public AI chatbots.",
        "It prevents accidental sharing of credentials, PII, source code and confidential data.",
        "It works directly inside the browser before information is transmitted.",
        "It supports browser-based AI platforms such as ChatGPT, Claude and Gemini.",
        "SENTINEL uses a multi-layer detection engine to identify sensitive information.",
        "More than 150 signature patterns are used to detect known sensitive data.",
        "Shannon Entropy is used to identify possible unknown secrets."
    ]
    create_bullet_slide("Abstract", s3_bullets, font_size=18, space_after=8)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 4: ABSTRACT (PART 2)
    # ═════════════════════════════════════════════════════════════════════════
    s4_bullets = [
        "The Luhn Algorithm is used to validate payment-card candidates.",
        "PII, passwords, credentials and confidential information are detected locally.",
        "Text normalization helps detect simple obfuscation techniques.",
        "High-risk prompts can be blocked before they are transmitted.",
        "Explainable warnings show users why their prompt was flagged.",
        "Auto-Redaction replaces sensitive values with safe placeholders.",
        "Masked security events are stored locally using SHA-256 hash chaining.",
        "Prompt data and telemetry are not sent to external servers.",
        "SENTINEL provides a lightweight and privacy-focused approach to AI data protection."
    ]
    create_bullet_slide("Abstract", s4_bullets, font_size=17, space_after=7)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 5: OBJECTIVES
    # ═════════════════════════════════════════════════════════════════════════
    s5_bullets = [
        "To detect sensitive information in AI prompts before it is transmitted to external AI services.",
        "To monitor user inputs and provide real-time security warnings.",
        "To identify API keys, passwords, PII, payment-card data and other confidential information.",
        "To calculate a risk score and classify detected information based on severity.",
        "To block high-risk prompts and prevent accidental data leakage.",
        "To provide automatic redaction so users can safely submit sanitized prompts."
    ]
    create_bullet_slide("Objectives", s5_bullets, font_size=19, space_after=12)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 6: LITERATURE SURVEY & GAPS (LITERATURE SURVEY)
    # ═════════════════════════════════════════════════════════════════════════
    s6_bullets = [
        "Traditional Data Loss Prevention (DLP) systems mainly use network monitoring, proxy servers and centralized security controls.",
        "Existing security approaches focus on preventing sensitive data from leaving organizational systems.",
        "AI chatbots introduce a new data-leakage risk because users directly enter information into external AI platforms.",
        "Existing solutions may not provide real-time, prompt-level protection directly inside the browser.",
        "There is a need for a lightweight solution specifically designed to inspect AI prompts before transmission."
    ]
    create_bullet_slide("Literature Survey & Gaps", s6_bullets, subheading="Literature Survey", font_size=18, space_after=10)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 7: RESEARCH GAPS
    # ═════════════════════════════════════════════════════════════════════════
    s7_bullets = [
        "Limited client-side protection specifically focused on AI chatbot prompts.",
        "Lack of immediate detection and blocking before sensitive data is transmitted.",
        "Limited user-friendly explanations about why a prompt is flagged.",
        "Lack of simple automatic redaction for detected sensitive information.",
        "Need for a privacy-focused approach that performs detection locally without sending prompt data to external servers."
    ]
    create_bullet_slide("Literature Survey & Gaps", s7_bullets, subheading="Research Gaps", font_size=18, space_after=12)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 8: EXISTING SOLUTIONS & LIMITATIONS (SOLUTIONS)
    # ═════════════════════════════════════════════════════════════════════════
    s8_bullets = [
        "Traditional DLP Systems: Monitor and control sensitive data moving through organizational networks.",
        "Network Proxies: Inspect network traffic and apply security policies before data reaches external services.",
        "Endpoint Security: Monitors user devices and applications to identify potentially sensitive activities.",
        "Organizational Policies: Depend on user awareness and rules to prevent employees from sharing confidential information.",
        "AI Usage Controls: Organizations can restrict or monitor access to external AI platforms."
    ]
    create_bullet_slide("Existing Solutions & Limitations", s8_bullets, subheading="Existing Solutions", font_size=18, space_after=10)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 9: EXISTING SOLUTIONS & LIMITATIONS (LIMITATIONS)
    # ═════════════════════════════════════════════════════════════════════════
    s9_bullets = [
        "These solutions may require additional infrastructure and configuration.",
        "Manual protection depends heavily on user awareness and careful behavior.",
        "Network-level monitoring may not provide direct prompt-level feedback.",
        "Users may not know which part of a prompt contains sensitive information.",
        "Existing approaches may not provide one-click automatic redaction before submission."
    ]
    create_bullet_slide("Existing Solutions & Limitations", s9_bullets, subheading="Limitations", font_size=18, space_after=12)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 10: PROPOSED SOLUTION
    # ═════════════════════════════════════════════════════════════════════════
    s10_bullets = [
        "SENTINEL is proposed as a client-side Data Loss Prevention browser extension for AI chatbots.",
        "It monitors user prompts and detects sensitive information before the request is transmitted.",
        "The detection engine uses 150+ signatures, Shannon Entropy, Luhn validation and PII detection.",
        "Detected information is classified according to its type, severity and associated risk score.",
        "High-risk prompts can be blocked, while sensitive values can be removed using Auto-Redaction.",
        "The system processes prompts locally and stores only masked audit information in browser storage."
    ]
    create_bullet_slide("Proposed Solution", s10_bullets, font_size=18, space_after=10)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 11: TECHNOLOGY STACK
    # ═════════════════════════════════════════════════════════════════════════
    s11_bullets = [
        "HTML5, CSS3 and JavaScript ES6+ are used to develop the browser extension.",
        "Chrome Manifest V3 provides the extension architecture and security model.",
        "Service Worker handles background processing and extension events.",
        "chrome.storage.local stores configuration and masked audit information locally.",
        "SHA-256 and Luhn Algorithm are used for audit security and payment-card validation.",
        "Node.js is used for testing and evaluating the detection system."
    ]
    create_bullet_slide("Technology Stack", s11_bullets, font_size=18, space_after=10)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 12: SYSTEM ARCHITECTURE FLOW (IMAGE)
    # ═════════════════════════════════════════════════════════════════════════
    s12 = prs.slides.add_slide(blank_slide_layout)
    if os.path.exists(diag_path):
        # The image already contains the title and logos, embed full slide perfectly
        s12.shapes.add_picture(diag_path, Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))
    else:
        add_logos(s12)
        add_title(s12, "SENTINEL – System Architecture Flow")

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 13: WORKING METHODOLOGY
    # ═════════════════════════════════════════════════════════════════════════
    s13_bullets = [
        "The user enters or pastes a prompt into a supported AI chatbot.",
        "SENTINEL monitors the input and intercepts the request before transmission.",
        "The prompt is normalized and analyzed using multiple detection techniques.",
        "Detected information is classified based on its type, severity and risk score.",
        "The system decides whether to Allow, Block or Redact the prompt.",
        "A masked security event is stored locally for auditing."
    ]
    create_bullet_slide("Working Methodology", s13_bullets, font_size=18, space_after=11)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 14: PROBLEM IMPACT (WHY IT MATTERS)
    # ═════════════════════════════════════════════════════════════════════════
    s14_bullets = [
        "Data Leakage → Sensitive information may be exposed to external AI services",
        "Credential Exposure → API keys and passwords may be misused",
        "Privacy Risk → Customer and personal information can be accidentally shared",
        "Intellectual Property Loss → Source code and confidential business data may be exposed",
        "Compliance Issues → Data exposure can lead to privacy and regulatory problems",
        "No Immediate Protection → Users may realize the risk only after submitting the prompt"
    ]
    create_bullet_slide("Problem Impact (Why it Matters)", s14_bullets, font_size=18, space_after=10)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 15: USE CASE SCENARIO
    # ═════════════════════════════════════════════════════════════════════════
    s15_bullets = [
        "Prompt Entry → Developer pastes authentication code into an AI chatbot for debugging.",
        "Sensitive Data → The code accidentally contains an API key and database credentials.",
        "Detection → SENTINEL scans the prompt and identifies the sensitive information.",
        "Warning & Block → A risk warning is displayed and the request is blocked if the risk is critical.",
        "Auto-Redaction → Sensitive values are replaced with safe placeholders such as [API_KEY].",
        "Safe Submission → The sanitized prompt is submitted to the AI chatbot without exposing the original credentials."
    ]
    create_bullet_slide("Use Case Scenario", s15_bullets, font_size=17, space_after=9)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 16: USER FLOW
    # ═════════════════════════════════════════════════════════════════════════
    s16_bullets = [
        "Open AI Chatbot → User opens a supported AI platform in the browser.",
        "Enter Prompt → User types or pastes a prompt into the chatbot.",
        "Monitor & Scan → SENTINEL monitors the input and scans for sensitive information.",
        "Risk Analysis → Detected data is classified and a risk score is calculated.",
        "Decision → Prompt is Allowed, Blocked or Redacted based on the risk level.",
        "Continue Safely → User submits the sanitized prompt, while a masked audit record is stored locally."
    ]
    create_bullet_slide("User Flow", s16_bullets, font_size=18, space_after=10)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 17: AI MODELS EXPLANATION
    # ═════════════════════════════════════════════════════════════════════════
    s17_bullets = [
        "Regex Model → Uses predefined patterns to detect known API keys, passwords, tokens and sensitive data.",
        "Shannon Entropy → Measures the randomness of text to identify possible unknown secrets.",
        "Luhn Algorithm → Validates possible payment-card numbers and reduces incorrect detections.",
        "PII Detection → Identifies personal information such as email addresses and phone numbers.",
        "TF-IDF + Logistic Regression → Experimental machine-learning model used to classify prompts based on their text features.",
        "Combined Detection → Multiple techniques work together to improve sensitive-data detection and risk assessment."
    ]
    create_bullet_slide("AI Models Explanation", s17_bullets, font_size=17, space_after=9)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 18: INNOVATION
    # ═════════════════════════════════════════════════════════════════════════
    s18_bullets = [
        "Pre-Transmission Protection → Detects sensitive information before the prompt reaches the external AI service.",
        "Multi-Layer Detection → Combines 150+ signatures, entropy, Luhn validation and PII detection.",
        "Privacy-First Processing → Performs prompt analysis locally without sending prompt data to external servers.",
        "Auto-Redaction → Automatically replaces detected sensitive values with safe placeholders.",
        "Explainable Alerts → Clearly shows users what type of sensitive information was detected and why.",
        "Tamper-Evident Audit → Uses SHA-256 hash chaining to maintain secure local audit records."
    ]
    create_bullet_slide("Innovation", s18_bullets, font_size=17, space_after=9)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 19: DATA FLOW (TECHNICAL)
    # ═════════════════════════════════════════════════════════════════════════
    s19_bullets = [
        "Prompt Input → User enters or pastes the prompt into the AI chatbot interface.",
        "Content Script → Monitors the webpage, input fields and user interactions.",
        "Request Interception → fetch and XMLHttpRequest requests are intercepted before transmission.",
        "Detection Engine → Prompt is normalized and scanned using Regex, Shannon Entropy, Luhn and PII detection.",
        "Decision Engine → Calculates the risk level and decides whether to Allow, Block or Redact the request.",
        "Local Storage → Masked audit information and configuration are stored using chrome.storage.local."
    ]
    create_bullet_slide("Data Flow (Technical)", s19_bullets, font_size=17, space_after=9)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 20: RESULTS & DISCUSSION
    # ═════════════════════════════════════════════════════════════════════════
    s20_bullets = [
        "High Detection Accuracy → The system achieved 100% accuracy on the evaluated test dataset.",
        "Reliable Detection → Precision and recall were both reported as 100%, with an F1 score of 1.00.",
        "Low False Positives → The reported false positive rate was 0%, reducing unnecessary security warnings.",
        "Fast Processing → Average detection latency was reported to be under 25 ms.",
        "Test Coverage → The evaluation included 95 prompts covering clean prompts, API keys, passwords, PII and confidential data.",
        "Overall Result → SENTINEL demonstrated effective real-time sensitive-data detection while maintaining local, privacy-focused processing."
    ]
    create_bullet_slide("Results & Discussion", s20_bullets, font_size=17, space_after=8)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 21: CHALLENGES & SOLUTIONS
    # ═════════════════════════════════════════════════════════════════════════
    s21_bullets = [
        "Dynamic AI Websites → MutationObserver detects changes in webpage elements and maintains monitoring.",
        "React-Controlled Inputs → Native setters and synthetic events help synchronize monitored input fields.",
        "Obfuscated Data → Text normalization handles unusual spacing, invisible characters and Unicode variations.",
        "Encoded Secrets → Base64 content is decoded and scanned again for sensitive information.",
        "False Positives → Entropy thresholds and Luhn validation help improve detection accuracy.",
        "Browser Limitations → The current system focuses on browser-based AI platforms rather than terminal applications."
    ]
    create_bullet_slide("Challenges & Solutions", s21_bullets, font_size=17, space_after=8)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 22: FUTURE SCOPE
    # ═════════════════════════════════════════════════════════════════════════
    s22_bullets = [
        "OCR Integration → Detect sensitive information present inside screenshots and images shared with AI platforms.",
        "SIEM Integration → Connect SENTINEL with enterprise security systems for centralized incident monitoring.",
        "Cross-Browser Support → Extend the solution to browsers such as Firefox and Safari.",
        "Advanced AI Detection → Improve the experimental machine-learning model for better prompt classification.",
        "More AI Platforms → Expand compatibility with additional AI chatbots and services.",
        "Enhanced Detection → Add more sensitive-data patterns and improve detection of complex obfuscation techniques."
    ]
    create_bullet_slide("Future Scope", s22_bullets, font_size=17, space_after=8)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 23: CONCLUSION
    # ═════════════════════════════════════════════════════════════════════════
    s23_bullets = [
        "Effective Protection → SENTINEL helps prevent sensitive information from being accidentally shared with public AI chatbots.",
        "Real-Time Detection → It detects and blocks high-risk prompts before they are transmitted.",
        "Privacy-Focused → Prompt analysis is performed locally without sending prompt data to external servers.",
        "Overall Outcome → SENTINEL provides a simple and practical solution for reducing AI-related data leakage."
    ]
    create_bullet_slide("Conclusion", s23_bullets, font_size=19, space_after=14)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 24: THANK YOU SLIDE (IMAGE)
    # ═════════════════════════════════════════════════════════════════════════
    s24 = prs.slides.add_slide(blank_slide_layout)
    if os.path.exists(thank_you_path):
        # Center the thank you graphic
        s24.shapes.add_picture(thank_you_path, Inches(2.9), Inches(1.2), width=Inches(7.5), height=Inches(5.1))
    else:
        add_title(s24, "Thank You!", top=Inches(3.0), font_size=48)

    for p in output_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        prs.save(p)
        print(f"Presentation saved successfully to: {p}")

if __name__ == "__main__":
    out1 = r"C:\Users\thero\Downloads\Mini_Project.pptx"
    out2 = r"c:\Users\thero\shadow-ai-prompt-guardian\Mini_Project.pptx"
    create_presentation([out1, out2])
